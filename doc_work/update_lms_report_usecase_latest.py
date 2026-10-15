import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


DOCX = Path(r"C:\DoAnCoSo\BaoCao_DoAnCoSoCNTT_NHOM_MMT.docx")
BACKUP = Path(r"C:\DoAnCoSo\LMS-Project\doc_work\BaoCao_DoAnCoSoCNTT_NHOM_MMT_truoc_sua_usecase_moi.docx")
USECASE_IMAGE = Path(r"C:\Users\ACER\Pictures\Screenshots\Screenshot 2026-06-10 090716.png")


ANALYSIS = [
    "Biểu đồ Use Case tổng quát mô tả phạm vi chức năng chính của hệ thống LMS Pro và mối quan hệ giữa các tác nhân với hệ thống. Hệ thống gồm ba tác nhân nghiệp vụ chính: Quản trị viên, Giảng viên và Học viên/Sinh viên. Ngoài ra, cơ sở dữ liệu MongoDB Atlas được thể hiện như hệ thống lưu trữ bên ngoài, nơi tiếp nhận và lưu trữ dữ liệu phát sinh từ các chức năng.",
    "Quản trị viên chịu trách nhiệm quản lý dữ liệu nền của hệ thống, bao gồm quản lý tài khoản người dùng, quản lý danh mục lớp học, chuyên ngành, môn học, quản lý phiên dạy, hỗ trợ tính lương tự động và import dữ liệu Excel. Các chức năng này giúp tạo dữ liệu ban đầu để giảng viên và sinh viên sử dụng trong quá trình học tập, điểm danh và theo dõi lịch sử.",
    "Giảng viên là tác nhân sử dụng các chức năng phục vụ quá trình giảng dạy. Giảng viên có thể xem lịch giảng dạy, quản lý lớp phụ trách, điểm danh thủ công, tạo mã QR điểm danh, quản lý buổi học, xem lịch sử điểm danh và xem lịch sử lương. Bên cạnh đó, hệ thống bổ sung nhóm chức năng Bài tập cho giảng viên, cho phép giảng viên giao bài tập cho từng lớp học, quản lý các bài tập đã giao, xem danh sách bài nộp của sinh viên, chấm điểm và nhập nhận xét cho bài làm.",
    "Sinh viên/Học viên sử dụng hệ thống để theo dõi quá trình học tập cá nhân. Sinh viên có thể xem tổng quan học tập, xem lớp học của tôi, xem lịch học, quét mã QR điểm danh và xem lịch sử điểm danh. Đối với chức năng Bài tập online, sinh viên có thể xem danh sách bài tập được giao, xem chi tiết yêu cầu bài tập, upload file để nộp bài, hủy nộp hoặc nộp lại nếu còn hạn, đồng thời xem điểm và nhận xét sau khi giảng viên chấm.",
    "Các mối quan hệ include/extend trong biểu đồ thể hiện sự phụ thuộc giữa các chức năng. Ví dụ, chức năng quên mật khẩu là luồng mở rộng của đăng nhập; chức năng tạo mã QR điểm danh có liên quan đến quản lý phiên dạy; chức năng nộp bài nằm trong nhóm bài tập online; chức năng chấm điểm và nhận xét gắn với quá trình xem bài nộp của sinh viên.",
]


def apply_body_font(paragraph, size=13, justify=True):
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(size)


def set_paragraph_text(paragraph, text, size=13, justify=True):
    for run in list(paragraph.runs):
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)
    apply_body_font(paragraph, size=size, justify=justify)


def find_paragraph_index(doc, predicate):
    for i, paragraph in enumerate(doc.paragraphs):
        if predicate(paragraph):
            return i
    raise ValueError("Target paragraph not found")


def replace_usecase_image(doc):
    caption_idx = find_paragraph_index(
        doc,
        lambda p: p.text.strip() == "Hình 3.1: Biểu đồ Use Case tổng quát của hệ thống LMS Pro",
    )
    image_idx = caption_idx - 1
    image_para = doc.paragraphs[image_idx]

    # Clear the old drawing from the paragraph and insert the new use case image.
    for run in list(image_para.runs):
        run._element.getparent().remove(run._element)
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.keep_with_next = True
    image_para.paragraph_format.space_after = Pt(3)
    image_para.add_run().add_picture(str(USECASE_IMAGE), width=Cm(16.0))

    caption = doc.paragraphs[caption_idx]
    set_paragraph_text(caption, "Hình 3.1: Biểu đồ Use Case tổng quát của hệ thống LMS Pro", size=13, justify=False)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = True
    return caption_idx


def update_usecase_analysis(doc, caption_idx):
    label_idx = caption_idx + 1
    set_paragraph_text(doc.paragraphs[label_idx], "Phân tích Hình 3.1:", size=13, justify=False)
    doc.paragraphs[label_idx].runs[0].font.bold = True
    doc.paragraphs[label_idx].paragraph_format.keep_with_next = True

    for offset, text in enumerate(ANALYSIS, start=2):
        paragraph = doc.paragraphs[caption_idx + offset]
        set_paragraph_text(paragraph, text, size=13, justify=True)


def normalize_assignment_bullets(doc):
    replacements = {
        "Quản lý bài tập đã giao.": "Quản lý danh sách bài tập đã giao.",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            set_paragraph_text(paragraph, replacements[text], size=13, justify=False)
            paragraph.style = "List Bullet"


def set_update_fields_on_open(doc):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        from docx.oxml import OxmlElement

        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    if not USECASE_IMAGE.exists():
        raise FileNotFoundError(USECASE_IMAGE)
    shutil.copyfile(DOCX, BACKUP)
    doc = Document(DOCX)
    caption_idx = replace_usecase_image(doc)
    update_usecase_analysis(doc, caption_idx)
    normalize_assignment_bullets(doc)
    set_update_fields_on_open(doc)
    doc.save(DOCX)
    print(DOCX)
    print(BACKUP)


if __name__ == "__main__":
    main()

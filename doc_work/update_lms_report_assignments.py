import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt


SOURCE = Path(r"C:\DoAnCoSo\BaoCao_DoAnCoSoCNTT_NHOM_MMT.docx")
OUTPUT = Path(r"C:\DoAnCoSo\BaoCao_DoAnCoSoCNTT_NHOM_MMT_BoSung_BaiTap.docx")


def fmt_paragraph(paragraph, justify=True):
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(13)


def insert_after(paragraph, text="", style=None, justify=True):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    fmt_paragraph(new_para, justify=justify)
    return new_para


def find_paragraph(doc, startswith=None, contains=None, exact=None):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if exact is not None and text == exact:
            return paragraph
        if startswith is not None and text.startswith(startswith):
            return paragraph
        if contains is not None and contains in text:
            return paragraph
    raise ValueError(f"Paragraph not found: startswith={startswith!r}, contains={contains!r}, exact={exact!r}")


def set_cell_text(cell, text, font_size=11):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(font_size)


def row_exists(table, first_cell_text):
    needle = first_cell_text.strip().lower()
    for row in table.rows:
        if row.cells and row.cells[0].text.strip().lower() == needle:
            return True
    return False


def append_row(table, values, font_size=11):
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value, font_size=font_size)
    return row


def set_update_fields_on_open(doc):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def update_tables(doc):
    routes = doc.tables[3]
    if not row_exists(routes, "/api/assignments"):
        append_row(
            routes,
            [
                "/api/assignments",
                "Nhóm API phục vụ chức năng bài tập online: tạo, lấy danh sách, xem chi tiết, cập nhật/xóa bài tập, nộp bài, hủy nộp và chấm điểm bài nộp.",
            ],
        )

    collections = doc.tables[4]
    if not row_exists(collections, "assignments"):
        append_row(
            collections,
            [
                "assignments",
                "title, description, classId, className, subjectName, instructorId, instructorName, dueDate, maxScore, attachments, status, createdAt, updatedAt",
                "Lưu thông tin bài tập do giảng viên giao cho từng lớp học, bao gồm yêu cầu bài tập, hạn nộp, điểm tối đa và trạng thái bài tập.",
            ],
            font_size=10,
        )
    if not row_exists(collections, "assignmentSubmissions"):
        append_row(
            collections,
            [
                "assignmentSubmissions",
                "assignmentId, studentId, studentName, classId, fileName, fileUrl, submittedAt, status, score, feedback, privateComment, createdAt, updatedAt",
                "Lưu thông tin bài nộp của sinh viên, file bài làm, trạng thái nộp bài, điểm số và nhận xét của giảng viên.",
            ],
            font_size=10,
        )

    tests = doc.tables[6]
    if not row_exists(tests, "Bài tập online"):
        append_row(
            tests,
            [
                "Bài tập online",
                "Giảng viên tạo bài tập, sinh viên xem chi tiết, upload file và nộp bài; giảng viên chấm điểm và nhập nhận xét.",
                "Bài tập hiển thị đúng theo lớp, trạng thái nộp được cập nhật, điểm và nhận xét hiển thị cho sinh viên sau khi chấm.",
            ],
        )

    appendix = doc.tables[7]
    endpoints = [
        ("POST", "/api/assignments", "Tạo bài tập."),
        ("GET", "/api/assignments/instructor/:instructorId", "Lấy bài tập theo giảng viên."),
        ("GET", "/api/assignments/student/:studentId", "Lấy bài tập của sinh viên."),
        ("GET", "/api/assignments/:id", "Xem chi tiết bài tập."),
        ("POST", "/api/assignments/:id/submit", "Sinh viên nộp bài."),
        ("DELETE", "/api/assignments/:id/unsubmit", "Sinh viên hủy nộp bài."),
        ("GET", "/api/assignments/:id/submissions", "Xem danh sách bài nộp."),
        ("PUT", "/api/assignments/submissions/:submissionId/grade", "Chấm điểm bài nộp."),
    ]
    existing = {(row.cells[0].text.strip(), row.cells[1].text.strip()) for row in appendix.rows if len(row.cells) >= 2}
    for method, endpoint, desc in endpoints:
        if (method, endpoint) not in existing:
            append_row(appendix, [method, endpoint, desc])


def update_api_section(doc):
    anchor = find_paragraph(doc, contains="endpoint điểm danh thủ công /api/attendance")
    p = insert_after(
        anchor,
        "Đối với chức năng bài tập online, báo cáo bổ sung nhóm API /api/assignments. Nhóm API này phục vụ việc tạo bài tập, lấy danh sách bài tập theo giảng viên hoặc sinh viên, xem chi tiết bài tập, cập nhật/xóa bài tập, nộp bài, hủy nộp bài và chấm điểm bài nộp.",
    )
    for item in [
        "GET /api/assignments/instructor/:instructorId - Lấy danh sách bài tập của giảng viên.",
        "GET /api/assignments/class/:classId - Lấy danh sách bài tập theo lớp.",
        "POST /api/assignments - Tạo bài tập mới.",
        "GET /api/assignments/:id - Xem chi tiết bài tập.",
        "PUT /api/assignments/:id - Cập nhật bài tập.",
        "DELETE /api/assignments/:id - Xóa bài tập.",
        "GET /api/assignments/:id/submissions - Lấy danh sách bài nộp của sinh viên.",
        "GET /api/assignments/student/:studentId - Lấy danh sách bài tập của sinh viên.",
        "POST /api/assignments/:id/submit - Sinh viên nộp bài.",
        "DELETE /api/assignments/:id/unsubmit - Sinh viên hủy nộp bài.",
        "PUT /api/assignments/submissions/:submissionId/grade - Giảng viên chấm điểm và nhận xét.",
    ]:
        p = insert_after(p, item, style="List Bullet", justify=False)


def update_feature_sections(doc):
    instructor_anchor = find_paragraph(doc, startswith="Lịch sử lương: ghi nhận thu nhập")
    p = insert_after(
        instructor_anchor,
        "Giảng viên được bổ sung chức năng Bài tập trên dashboard. Chức năng này cho phép giảng viên tạo bài tập cho các lớp học mình phụ trách, bao gồm các thông tin như tiêu đề bài tập, mô tả yêu cầu, lớp học, môn học, hạn nộp, điểm tối đa và trạng thái bài tập. Sau khi giao bài, giảng viên có thể xem danh sách bài tập đã tạo, mở chi tiết từng bài tập, theo dõi danh sách sinh viên đã nộp hoặc chưa nộp bài, xem file bài làm của sinh viên, nhập điểm và nhận xét cho từng bài nộp. Việc bổ sung chức năng này giúp hệ thống mở rộng từ quản lý lớp học và điểm danh sang hỗ trợ quá trình học tập trực tuyến.",
    )
    for item in [
        "Tạo bài tập theo lớp học.",
        "Theo dõi trạng thái nộp bài của sinh viên.",
        "Xem file bài nộp.",
        "Chấm điểm và nhận xét bài làm.",
        "Quản lý bài tập đã giao.",
    ]:
        p = insert_after(p, item, style="List Bullet", justify=False)

    student_anchor = find_paragraph(doc, startswith="Thông tin cá nhân: cập nhật hồ sơ")
    p = insert_after(
        student_anchor,
        "Sinh viên được bổ sung chức năng Bài tập online trên dashboard. Chức năng này cho phép sinh viên xem danh sách bài tập được giao theo các lớp đang học, lọc bài tập theo trạng thái như chưa nộp, đã nộp, quá hạn hoặc đã chấm. Khi mở chi tiết bài tập, sinh viên có thể xem yêu cầu bài tập, hạn nộp, điểm tối đa và các thông tin liên quan. Sinh viên có thể upload file bài làm và nộp bài trực tiếp trên hệ thống. Nếu bài tập còn hạn, sinh viên có thể hủy nộp hoặc nộp lại bài. Sau khi giảng viên chấm, sinh viên có thể xem điểm và nhận xét. Chức năng này giúp quá trình nhận bài, nộp bài và theo dõi kết quả bài tập được tập trung trong hệ thống LMS.",
    )
    for item in [
        "Xem danh sách bài tập được giao.",
        "Xem chi tiết yêu cầu bài tập.",
        "Upload file và nộp bài trực tuyến.",
        "Hủy nộp hoặc nộp lại nếu còn hạn.",
        "Xem điểm và nhận xét sau khi giảng viên chấm.",
    ]:
        p = insert_after(p, item, style="List Bullet", justify=False)


def update_chapter4(doc):
    instructor_anchor = find_paragraph(doc, startswith="Đối với vai trò giảng viên, các chức năng tập trung")
    insert_after(
        instructor_anchor,
        "Ở phiên bản cập nhật, dashboard giảng viên được bổ sung chức năng Bài tập. Chức năng này cho phép giảng viên tạo bài tập cho từng lớp học, nhập mô tả yêu cầu, hạn nộp và điểm tối đa. Giảng viên có thể xem danh sách bài tập đã giao, mở chi tiết bài tập để theo dõi trạng thái nộp bài của sinh viên, xem file bài làm, nhập điểm và nhận xét. Kết quả thử nghiệm cho thấy chức năng này hỗ trợ giảng viên quản lý hoạt động học tập online thuận tiện hơn, thay vì chỉ quản lý lớp học và điểm danh.",
    )

    student_anchor = find_paragraph(doc, startswith="Đối với vai trò sinh viên, hệ thống tập trung")
    insert_after(
        student_anchor,
        "Đối với vai trò sinh viên, hệ thống được bổ sung chức năng Bài tập online. Sinh viên có thể xem danh sách bài tập được giao theo lớp, theo dõi hạn nộp, xem yêu cầu chi tiết và nộp file bài làm trực tiếp trên hệ thống. Sau khi nộp bài, hệ thống hiển thị trạng thái đã nộp; sinh viên có thể theo dõi điểm số và nhận xét sau khi giảng viên chấm. Chức năng này giúp sinh viên quản lý bài tập tập trung hơn và giảm phụ thuộc vào việc nộp bài qua các kênh rời rạc.",
    )


def replace_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)
    fmt_paragraph(paragraph, justify=paragraph.style.name != "List Bullet")


def update_chapter5(doc):
    anchor = find_paragraph(doc, startswith="Có cơ chế OTP qua email")
    insert_after(
        anchor,
        "Hệ thống đã bổ sung chức năng bài tập online, cho phép giảng viên giao bài tập và sinh viên nộp bài trực tuyến. Chức năng này giúp hệ thống LMS hoàn thiện hơn vì không chỉ hỗ trợ quản lý lớp học và điểm danh, mà còn hỗ trợ thêm hoạt động học tập trực tuyến.",
        style="List Bullet",
        justify=False,
    )

    old_limit = find_paragraph(doc, startswith="Chưa có module nội dung học tập như")
    replace_paragraph_text(
        old_limit,
        "Chưa có module nội dung học tập nâng cao như bài giảng, tài liệu học tập, điểm số tổng hợp hoặc thông báo realtime.",
    )
    insert_after(
        old_limit,
        "Chức năng bài tập hiện mới dừng ở mức nộp file và chấm điểm cơ bản, chưa hỗ trợ bình luận realtime, thông báo nhắc hạn nộp hoặc kiểm tra trùng lặp nội dung bài làm.",
        style="List Bullet",
        justify=False,
    )

    old_direction = find_paragraph(doc, startswith="Bổ sung chức năng tài liệu học tập, bài tập")
    replace_paragraph_text(
        old_direction,
        "Có thể phát triển thêm thông báo bài tập, nhắc hạn nộp, bình luận theo từng bài tập, thống kê tiến độ nộp bài và xuất báo cáo điểm bài tập.",
    )

    conclusion = find_paragraph(doc, startswith="Đề tài E-Learning LMS đã xây dựng được nền tảng")
    replace_paragraph_text(
        conclusion,
        "Đề tài E-Learning LMS đã xây dựng được nền tảng quản lý học tập ở mức cốt lõi với kiến trúc fullstack, giao diện theo vai trò và các chức năng nghiệp vụ quan trọng. Admin có thể quản lý dữ liệu nền, giảng viên có thể quản lý lớp, điểm danh và giao bài tập, sinh viên có thể theo dõi lịch học, quét QR và nộp bài tập online. Việc áp dụng Next.js, React, TypeScript, Tailwind CSS, Node.js, Express và MongoDB giúp project thể hiện đầy đủ quy trình xây dựng một ứng dụng LMS hiện đại ở quy mô đồ án cơ sở.",
    )


def main():
    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    update_tables(doc)
    update_api_section(doc)
    update_feature_sections(doc)
    update_chapter4(doc)
    update_chapter5(doc)
    set_update_fields_on_open(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

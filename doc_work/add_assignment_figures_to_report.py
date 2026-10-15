import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt


DOCX = Path(r"C:\DoAnCoSo\BaoCao_DoAnCoSoCNTT_NHOM_MMT.docx")
BACKUP = Path(r"C:\DoAnCoSo\LMS-Project\doc_work\BaoCao_DoAnCoSoCNTT_NHOM_MMT_truoc_them_hinh_baitap.docx")

FIGURES = [
    (
        Path(r"C:\Users\ACER\Pictures\Screenshots\Screenshot 2026-06-10 100025.png"),
        "Hình 4.19: Dashboard Giảng viên - chức năng bài tập",
    ),
    (
        Path(r"C:\Users\ACER\Pictures\Screenshots\Screenshot 2026-06-10 100230.png"),
        "Hình 4.20: Dashboard Giảng viên - chi tiết bài tập và bài nộp sinh viên",
    ),
    (
        Path(r"C:\Users\ACER\Pictures\Screenshots\Screenshot 2026-06-10 100102.png"),
        "Hình 4.21: Dashboard Sinh viên - bài tập online",
    ),
    (
        Path(r"C:\Users\ACER\Pictures\Screenshots\Screenshot 2026-06-10 100204.png"),
        "Hình 4.22: Dashboard Sinh viên - chi tiết bài tập và nộp bài",
    ),
]

TABLE_41_ROWS = [
    (
        "Hình 4.19",
        "Dashboard Giảng viên - chức năng bài tập",
        "Minh họa giao diện tạo và quản lý bài tập do giảng viên giao cho sinh viên theo từng lớp học.",
    ),
    (
        "Hình 4.20",
        "Dashboard Giảng viên - chi tiết bài tập và bài nộp sinh viên",
        "Minh họa danh sách sinh viên đã nộp/chưa nộp, file bài làm, điểm và nhận xét.",
    ),
    (
        "Hình 4.21",
        "Dashboard Sinh viên - bài tập online",
        "Minh họa danh sách bài tập được giao cho sinh viên theo các lớp đang học.",
    ),
    (
        "Hình 4.22",
        "Dashboard Sinh viên - chi tiết bài tập và nộp bài",
        "Minh họa giao diện xem yêu cầu bài tập, hạn nộp, upload file, nộp bài và theo dõi trạng thái.",
    ),
]

TEACHER_DESC = (
    "Dashboard giảng viên được bổ sung chức năng Bài tập. Chức năng này cho phép giảng viên tạo bài tập theo từng lớp học, "
    "nhập mô tả yêu cầu, hạn nộp và điểm tối đa. Giảng viên có thể xem danh sách bài tập đã giao, theo dõi trạng thái nộp bài "
    "của sinh viên, xem file bài làm, nhập điểm và nhận xét."
)

STUDENT_DESC = (
    "Dashboard sinh viên được bổ sung chức năng Bài tập online. Sinh viên có thể xem danh sách bài tập được giao, mở chi tiết "
    "bài tập, theo dõi hạn nộp, upload file bài làm và nộp trực tiếp trên hệ thống. Sau khi giảng viên chấm, sinh viên có thể "
    "xem điểm và nhận xét."
)


def format_run(run, size=13):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def format_para(paragraph, justify=True, size=13):
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        format_run(run, size)


def set_paragraph_text(paragraph, text, style_name=None, justify=True):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)
    if style_name:
        paragraph.style = style_name
    format_para(paragraph, justify=justify)


def insert_paragraph_before(target, style_name=None):
    new_p = OxmlElement("w:p")
    target._p.addprevious(new_p)
    para = Paragraph(new_p, target._parent)
    if style_name:
        para.style = style_name
    return para


def add_tc_field(paragraph, caption_text):
    # Word uses TC fields in this report's figure list, so add one for each new caption.
    def add_run_with(child):
        run = OxmlElement("w:r")
        run.append(child)
        paragraph._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    add_run_with(begin)

    instr1 = OxmlElement("w:instrText")
    instr1.set(qn("xml:space"), "preserve")
    instr1.text = ' TC "'
    add_run_with(instr1)

    instr2 = OxmlElement("w:instrText")
    instr2.text = caption_text
    add_run_with(instr2)

    instr3 = OxmlElement("w:instrText")
    instr3.set(qn("xml:space"), "preserve")
    instr3.text = '" \\f F '
    add_run_with(instr3)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    add_run_with(end)


def add_figure_before(target, image_path, caption_text):
    image_para = insert_paragraph_before(target, "Normal")
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.keep_with_next = True
    image_para.paragraph_format.space_after = Pt(3)
    image_para.add_run().add_picture(str(image_path), width=Cm(16.0))

    caption_para = insert_paragraph_before(target, "Caption")
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_para.paragraph_format.keep_together = True
    caption_para.paragraph_format.keep_with_next = True
    run = caption_para.add_run(caption_text)
    format_run(run, 13)
    add_tc_field(caption_para, caption_text)
    return caption_para


def table_has_figure(table, fig_no):
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == fig_no:
            return True
    return False


def set_cell_text(cell, text, size=10.5):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            format_run(run, size)


def update_table_41(doc):
    table = doc.tables[5]
    for values in TABLE_41_ROWS:
        if table_has_figure(table, values[0]):
            continue
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)


def remove_existing_new_figures(doc):
    # Idempotency guard: if the script is run twice, remove existing 4.19-4.22 image+caption blocks first.
    labels = {caption for _, caption in FIGURES}
    for i in range(len(doc.paragraphs) - 1, -1, -1):
        paragraph = doc.paragraphs[i]
        if paragraph.text.strip() in labels:
            prev = doc.paragraphs[i - 1] if i > 0 else None
            paragraph._p.getparent().remove(paragraph._p)
            if prev is not None and prev._p.xpath(".//pic:pic"):
                prev._p.getparent().remove(prev._p)


def main():
    missing = [str(path) for path, _ in FIGURES if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing figure images: " + ", ".join(missing))

    shutil.copyfile(DOCX, BACKUP)
    doc = Document(DOCX)

    # Update concise descriptions around the new images.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Ở phiên bản cập nhật, dashboard giảng viên được bổ sung chức năng Bài tập"):
            set_paragraph_text(paragraph, TEACHER_DESC, "Normal", justify=True)
        elif text.startswith("Đối với vai trò sinh viên, hệ thống được bổ sung chức năng Bài tập online"):
            set_paragraph_text(paragraph, STUDENT_DESC, "Normal", justify=True)

    remove_existing_new_figures(doc)

    target = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() == "4.2 Ý nghĩa các màn hình minh họa"
    )
    for image_path, caption in FIGURES:
        add_figure_before(target, image_path, caption)

    update_table_41(doc)
    doc.save(DOCX)
    print(DOCX)
    print(BACKUP)


if __name__ == "__main__":
    main()
